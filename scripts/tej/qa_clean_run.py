"""A4/A5 verification, test 3: clean reproducibility run (reference = Run 4, the canonical set after A5).

Builds a fresh copy of the DER inputs in DER_TEJ/QA_clean, deletes every output the chain produces,
regenerates C2 with the fixed builder, runs the whole province chain, master results and tables,
then the TEJ steps, figures and documents in a fresh TEJ workspace, and compares everything with Run 4.

DER and the Run 3 mirror are only read. Usage: python qa_clean_run.py [setup|run|compare|all]
"""
import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

PY = r"C:\Users\user\AppData\Local\Microsoft\WindowsApps\python3.exe"
E = Path(r"D:\Clude Code GIS\Energy")
DER = E / "DER"
REPO = E / "DER_paper_repo"                       # branch fix/c2-poverty-join (fixed builder) must be checked out
TEJ = E / "DER_TEJ"
REF = TEJ / "Rerun_C2fix" / "Energy" / "DER"
Q = TEJ / "QA_clean"
QE = Q / "Energy"
QD = QE / "DER"
QT = Q / "DER_TEJ"
PROC_OUT = ["C2_poverty_province.csv", "C2_poverty_90m.tif", "province_PAI_summary.csv", "EWM_weights.csv",
            "province_criteria_normalised.csv", "province_LISA.csv", "province_LISA_fdr.csv",
            "reviewer_robustness_summary.json", "sensitivity_analysis.csv", "sensitivity_summary.csv",
            "bootstrap_uncertainty.json", "validation_solar_osm.csv", "validation_zone_distribution.csv",
            "validation_zone_stats.json", "spatial_permutation_results.json", "maup_results.json",
            "sensitivity_pai_mv.csv", "sensitivity_pai_mv.json", "validation_ser_external.csv",
            "cluster_stability.json", "lisa_fdr.json", "criterion_scaling.json", "C2_poverty_join_diagnostics.txt"]
CHAIN = ["02_Preprocessing/build_c2_poverty_province.py", "02_Preprocessing/rasterize_c2_poverty.py",
         "06_Priority_Areas/compute_pai_province_level.py", "06_Priority_Areas/spatial_autocorrelation.py",
         "06_Priority_Areas/reviewer_robustness_checks.py", "06_Priority_Areas/sensitivity_analysis.py",
         "06_Priority_Areas/bootstrap_weights.py", "06_Priority_Areas/validate_solar_osm.py",
         "06_Priority_Areas/sensitivity_validation_province.py", "06_Priority_Areas/validate_zones_province.py",
         "06_Priority_Areas/spatial_permutation_test.py", "06_Priority_Areas/maup_analysis.py",
         "06_Priority_Areas/sensitivity_pai_mv.py", "06_Priority_Areas/validate_ser_external.py"]
TEJ_STEPS = ["province_metrics.py", "pai_grid_qa.py", "scenarios.py", "zone_distances.py", "zone_policy_table.py",
             "erp_posthoc_diagnostic.py", "area_weight_check.py",
             "figures/fig01_framework_tej.py", "figures/fig02_study_area_tej.py", "figures/fig03_input_layers_tej.py",
             "figures/fig06_07_pai_zone_maps_tej.py", "figures/fig10_deployment_scatter_tej.py",
             "figures/fig13_14_radar_policy_tej.py", "figures/make_supp_fig_scripts_tej.py",
             "figures/fig08_lisa_map_tej.py", "figures/fig11_sensitivity_tej.py", "figures/fig12_bangkok_district_tej.py", "figures/fig_s13ab_corr_weights_tej.py",
             "assemble_manuscript_tej.py", "assemble_supplementary_tej.py", "export_figures_tej.py", "figure_manifest.py"]
MIRROR_TXT = r"D:\Clude Code GIS\Energy\DER_TEJ\Rerun_C2fix\Energy"


def sh(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def setup():
    if Q.exists():
        shutil.rmtree(Q)
    (QE / "Energy map" / "GridNetwork_OSM").mkdir(parents=True)
    for d in ("Data", "Tables"):
        shutil.copytree(DER / d, QD / d)
    shutil.copytree(REPO / "Scripts", QD / "Scripts", ignore=shutil.ignore_patterns("__pycache__"))
    (QD / "Figures").mkdir()
    for f in ("Thailand_DEM_90m.tif", "สถิติความยากจนและการกระจายรายได้_2567.xlsx"):
        shutil.copyfile(E / "Energy map" / f, QE / "Energy map" / f)
    shutil.copyfile(E / "Energy map" / "GridNetwork_OSM" / "plants.geojson", QE / "Energy map" / "GridNetwork_OSM" / "plants.geojson")
    shutil.copytree(E / "Energy map" / "Thailand_Provincial_Socioeconomic_Resilience_Index_2566_2567",
                    QE / "Energy map" / "Thailand_Provincial_Socioeconomic_Resilience_Index_2566_2567")
    # same two mirror-only adjustments as Run 3: SciPy chi-square on rounded counts, absolute paths
    vz = QD / "Scripts" / "06_Priority_Areas" / "validate_zones_province.py"
    s = vz.read_text(encoding="utf-8")
    s = s.replace("chi2, p = stats.chisquare(D.solar_obs, D.expected)",
                  "chi2 = float((((D.solar_obs - D.expected) ** 2) / D.expected).sum())\np = float(stats.chi2.sf(chi2, len(D) - 1))")
    vz.write_text(s, encoding="utf-8")
    for p in (QD / "Scripts").rglob("*.py"):
        t = p.read_text(encoding="utf-8", errors="surrogateescape")
        o = t
        t = t.replace("D:\\Clude Code GIS\\Energy", str(QE)).replace("D:/Clude Code GIS/Energy", str(QE).replace("\\", "/"))
        if t != o:
            p.write_text(t, encoding="utf-8", errors="surrogateescape")
    # empty output state: every chain output removed
    for f in PROC_OUT:
        (QD / "Data" / "14_Processed_90m" / f).unlink(missing_ok=True)
    (QD / "Data" / "master_results.json").unlink(missing_ok=True)
    (QD / "Tables" / "DER_Paper_Tables.xlsx").unlink(missing_ok=True)
    # fresh TEJ workspace: scripts + unchanged inputs only
    QT.mkdir(parents=True)
    shutil.copytree(TEJ / "Scripts", QT / "Scripts", ignore=shutil.ignore_patterns("__pycache__"))
    for d in ("mv_distance", "policy"):
        shutil.copytree(TEJ / "Data" / d, QT / "Data" / d)
    shutil.copyfile(TEJ / "Data" / "references.json", QT / "Data" / "references.json")
    # the manuscript counts protocol amendments from the protocol itself (review 2026-10-18)
    (QT / "Docs").mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEJ / "Docs" / "SCENARIO_PROTOCOL.md", QT / "Docs" / "SCENARIO_PROTOCOL.md")
    for f in ("pea_linkage_qa.json",):
        if (TEJ / "Data" / "qa" / f).exists():
            (QT / "Data" / "qa").mkdir(parents=True, exist_ok=True)
            shutil.copyfile(TEJ / "Data" / "qa" / f, QT / "Data" / "qa" / f)
    print("setup done")


def run_step(script, cwd, env):
    r = subprocess.run([PY, str(script)], cwd=str(cwd), env=env, capture_output=True, text=True, encoding="utf-8", errors="replace")
    print(f"{script.name}: rc={r.returncode}", flush=True)
    if r.returncode != 0:
        print(r.stdout[-1500:], r.stderr[-2500:])
        sys.exit(1)


def run():
    env = dict(os.environ, ENERGY_DATA_ROOT=str(QE), DER_PROJECT_ROOT=str(QD), PYTHONIOENCODING="utf-8", MPLBACKEND="Agg")
    for s in CHAIN:
        run_step(QD / "Scripts" / s, (QD / "Scripts" / s).parent, env)
    r = subprocess.run([PY, str(TEJ / "Scripts" / "guard_c2_scaling.py"), str(QD)], env=env, capture_output=True, text=True, encoding="utf-8", errors="replace")
    print("guard_c2_scaling.py: rc=", r.returncode, flush=True)
    assert r.returncode == 0, r.stdout[-2000:]
    run_step_args = [PY, str(TEJ / "Scripts" / "rerun_c2fix_extras.py"), str(QD / "Data" / "14_Processed_90m")]
    r = subprocess.run(run_step_args, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace")
    print("rerun_c2fix_extras.py: rc=", r.returncode, flush=True)
    assert r.returncode == 0, r.stderr[-1500:]
    run_step(QD / "Scripts" / "09_Manuscript" / "build_master_results.py", QD / "Scripts" / "09_Manuscript", env)
    run_step(QD / "Scripts" / "08_Tables" / "generate_tables.py", QD / "Scripts" / "08_Tables", env)
    tenv = dict(env, DER_SOURCE_ROOT=str(QD))
    for s in TEJ_STEPS:
        run_step(QT / "Scripts" / s, (QT / "Scripts" / s).parent, tenv)
    print("clean run done")


def compare():
    import docx
    import numpy as np
    import pandas as pd
    from PIL import Image
    rep = {"der_outputs": {}, "tej_outputs": {}, "figures": {}, "docx_text": {}}
    for f in PROC_OUT:
        a, b = QD / "Data" / "14_Processed_90m" / f, REF / "Data" / "14_Processed_90m" / f
        if not a.exists() or not b.exists():
            rep["der_outputs"][f] = "missing in " + ("clean" if not a.exists() else "run4"); continue
        if sh(a) == sh(b):
            rep["der_outputs"][f] = "identical"
        elif f.endswith(".csv"):
            x, y = pd.read_csv(a), pd.read_csv(b)
            same = x.shape == y.shape and all(
                (np.allclose(x[c], y[c], atol=1e-9, equal_nan=True) if pd.api.types.is_numeric_dtype(x[c]) else x[c].astype(str).equals(y[c].astype(str)))
                for c in x.columns if c in y.columns) and list(x.columns) == list(y.columns)
            rep["der_outputs"][f] = "same values (file bytes differ)" if same else "DIFFERENT"
        elif f.endswith(".txt"):
            rep["der_outputs"][f] = "DIFFERENT"
        elif f.endswith(".json"):
            rep["der_outputs"][f] = "same values (formatting differs)" if json.loads(a.read_text(encoding="utf-8")) == json.loads(b.read_text(encoding="utf-8")) else "DIFFERENT"
        else:
            import rasterio
            with rasterio.open(a) as u, rasterio.open(b) as v:
                rep["der_outputs"][f] = "same values" if np.array_equal(u.read(1), v.read(1), equal_nan=True) else "DIFFERENT"
    ma = json.loads((QD / "Data" / "master_results.json").read_text(encoding="utf-8"))
    mb = json.loads((REF / "Data" / "master_results.json").read_text(encoding="utf-8"))
    rep["der_outputs"]["master_results.json"] = "same values" if ma == mb else "DIFFERENT"
    xa = pd.read_excel(QD / "Tables" / "DER_Paper_Tables.xlsx", sheet_name=None)
    xb = pd.read_excel(REF / "Tables" / "DER_Paper_Tables.xlsx", sheet_name=None)
    rep["der_outputs"]["DER_Paper_Tables.xlsx"] = "same values" if xa.keys() == xb.keys() and all(xa[k].astype(str).equals(xb[k].astype(str)) for k in xa) else "DIFFERENT"
    snap = json.loads((TEJ / "Data" / "qa" / "F_snapshot.json").read_text(encoding="utf-8"))["files"]
    for k in snap:
        if k.startswith("Rerun_C2fix"):
            continue
        # Run 1 (superseded by run 2, amendment A2) is kept as a record only, no chain script writes it
        # (Docs/SCENARIO_RUN_LOG.md). Excluded 2026-10-18 after the clean run reported it missing.
        if k.replace("\\", "/").endswith("run1_scenario_tradeoff_noMV_77prov.csv"):
            rep["tej_outputs"][k] = "record only (run 1, superseded), not regenerated"
            continue
        a = QT / k
        rep["tej_outputs"][k] = ("identical" if a.exists() and sh(a) == snap[k] else ("missing" if not a.exists() else "DIFFERENT"))
    extra = ["Data/erp_diagnostic/erp_diagnostic.json", "Data/qa/pai_grid_qa.json"]
    for k in extra:
        a, b = QT / k, TEJ / k
        rep["tej_outputs"][k] = "identical" if a.exists() and sh(a) == sh(b) else "DIFFERENT"
    for p in sorted(list((TEJ / "Figures" / "main").glob("*.png")) + list((TEJ / "Figures" / "supplementary").glob("*.png"))):
        q = QT / "Figures" / p.parent.name / p.name
        if not q.exists():
            rep["figures"][p.name] = "missing"; continue
        u, v = np.asarray(Image.open(p)), np.asarray(Image.open(q))
        rep["figures"][p.name] = "identical pixels" if u.shape == v.shape and np.array_equal(u, v) else "DIFFERENT"
    for name in ("DER_TEJ_Manuscript", "DER_TEJ_Supplementary", "DER_TEJ_Highlights"):
        def text(path):
            d = docx.Document(str(path))
            return "\n".join(x.text for x in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        rep["docx_text"][name] = "identical text" if text(TEJ / "Manuscript" / f"{name}.docx") == text(QT / "Manuscript" / f"{name}.docx") else "DIFFERENT"
    out = TEJ / "Data" / "qa" / "clean_run_compare.json"
    out.write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    bad = {g: {k: v for k, v in d.items() if "DIFF" in v or "missing" in v} for g, d in rep.items()}
    print(json.dumps(bad, ensure_ascii=False, indent=1))
    print("checked:", {g: len(d) for g, d in rep.items()})


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "all"
    if mode in ("setup", "all"):
        setup()
    if mode in ("run", "all"):
        run()
    if mode in ("compare", "all"):
        compare()
