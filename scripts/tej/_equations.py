"""Editable OMML equations shared by the assemblers (moved from the main assembler, F step 2a)."""
import re
from lxml import etree
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

_STATE = {'doc': None, 'font': None, 'prefix': ''}


def init_eq(doc, font_fn, prefix=''):
    _STATE.update(doc=doc, font=font_fn, prefix=prefix)


_MML_NS = "http://www.w3.org/1998/Math/MathML"
_MML2OMML = etree.XSLT(etree.parse(
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
))

_EQUATION_MATHML = {
    1: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msub><mi>r</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>=</mo>
      <mfrac><mrow><msub><mi>x</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>-</mo><msub><mi>x</mi><mrow><mi>min</mi><mo>,</mo><mi>j</mi></mrow></msub></mrow>
      <mrow><msub><mi>x</mi><mrow><mi>max</mi><mo>,</mo><mi>j</mi></mrow></msub><mo>-</mo><msub><mi>x</mi><mrow><mi>min</mi><mo>,</mo><mi>j</mi></mrow></msub></mrow></mfrac>
    </mrow></math>""",
    2: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msub><mi>p</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>=</mo>
      <mfrac><msub><mi>r</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub>
      <mrow><munder><mo largeop="true">&#x2211;</mo><mi>i</mi></munder><msub><mi>r</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub></mrow></mfrac>
    </mrow></math>""",
    3: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msub><mi>e</mi><mi>j</mi></msub><mo>=</mo><mo>-</mo>
      <mfrac><mn>1</mn><mrow><mi>ln</mi><mo>(</mo><mi>m</mi><mo>)</mo></mrow></mfrac><mo>&#x22c5;</mo>
      <munder><mo largeop="true">&#x2211;</mo><mi>i</mi></munder>
      <mrow><mo>(</mo><msub><mi>p</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>&#x22c5;</mo><mi>ln</mi><msub><mi>p</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>)</mo></mrow>
      <mo>,</mo><mspace width="0.5em"/><mtext>with </mtext><msub><mi>p</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mi>ln</mi><msub><mi>p</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>=</mo><mn>0</mn><mtext> when </mtext><msub><mi>p</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>=</mo><mn>0</mn>
    </mrow></math>""",
    4: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>d</mi><mi>j</mi></msub><mo>=</mo><mn>1</mn><mo>-</mo><msub><mi>e</mi><mi>j</mi></msub></mrow></math>""",
    5: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msub><mi>w</mi><mi>j</mi></msub><mo>=</mo><mfrac><msub><mi>d</mi><mi>j</mi></msub><mrow><munder><mo largeop="true">&#x2211;</mo><mi>j</mi></munder><msub><mi>d</mi><mi>j</mi></msub></mrow></mfrac>
      <mo>,</mo><mspace width="0.8em"/><munder><mo largeop="true">&#x2211;</mo><mi>j</mi></munder><msub><mi>w</mi><mi>j</mi></msub><mo>=</mo><mn>1</mn>
    </mrow></math>""",
    6: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msup><mi>A</mi><mo>+</mo></msup><mo>=</mo><mo>{</mo><msubsup><mi>v</mi><mi>j</mi><mo>+</mo></msubsup><mo>=</mo><msub><mi>w</mi><mi>j</mi></msub><mo>&#x22c5;</mo><msubsup><mi>r</mi><mi>j</mi><mo>+</mo></msubsup><mo>=</mo><msub><mi>w</mi><mi>j</mi></msub><mo>&#x22c5;</mo><mn>1</mn><mo>=</mo><msub><mi>w</mi><mi>j</mi></msub><mo>|</mo><mi>j</mi><mo>&#x2208;</mo><mtext>benefit</mtext><mo>}</mo>
    </mrow></math>""",
    7: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msup><mi>A</mi><mo>-</mo></msup><mo>=</mo><mo>{</mo><msubsup><mi>v</mi><mi>j</mi><mo>-</mo></msubsup><mo>=</mo><msub><mi>w</mi><mi>j</mi></msub><mo>&#x22c5;</mo><msubsup><mi>r</mi><mi>j</mi><mo>-</mo></msubsup><mo>=</mo><msub><mi>w</mi><mi>j</mi></msub><mo>&#x22c5;</mo><mn>0</mn><mo>=</mo><mn>0</mn><mo>|</mo><mi>j</mi><mo>&#x2208;</mo><mtext>benefit</mtext><mo>}</mo>
    </mrow></math>""",
    8: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msubsup><mi>d</mi><mi>i</mi><mo>+</mo></msubsup><mo>=</mo><msqrt><mrow><munder><mo largeop="true">&#x2211;</mo><mi>j</mi></munder><msup><mrow><mo>(</mo><msub><mi>v</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>-</mo><msubsup><mi>v</mi><mi>j</mi><mo>+</mo></msubsup><mo>)</mo></mrow><mn>2</mn></msup></mrow></msqrt>
    </mrow></math>""",
    9: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msubsup><mi>d</mi><mi>i</mi><mo>-</mo></msubsup><mo>=</mo><msqrt><mrow><munder><mo largeop="true">&#x2211;</mo><mi>j</mi></munder><msup><mrow><mo>(</mo><msub><mi>v</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>-</mo><msubsup><mi>v</mi><mi>j</mi><mo>-</mo></msubsup><mo>)</mo></mrow><mn>2</mn></msup></mrow></msqrt>
    </mrow></math>""",
    10: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <msub><mi>PAI</mi><mi>i</mi></msub><mo>=</mo><mfrac><msubsup><mi>d</mi><mi>i</mi><mo>-</mo></msubsup><mrow><msubsup><mi>d</mi><mi>i</mi><mo>+</mo></msubsup><mo>+</mo><msubsup><mi>d</mi><mi>i</mi><mo>-</mo></msubsup></mrow></mfrac><mo>,</mo><mspace width="0.8em"/><mi>PAI</mi><mo>&#x2208;</mo><mo>[</mo><mn>0</mn><mo>,</mo><mn>1</mn><mo>]</mo>
    </mrow></math>""",
    11: """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>
      <mi>I</mi><mo>=</mo><mfrac><mi>n</mi><msub><mi>S</mi><mn>0</mn></msub></mfrac><mo>&#x22c5;</mo>
      <mfrac><mrow><munder><mo largeop="true">&#x2211;</mo><mi>i</mi></munder><munder><mo largeop="true">&#x2211;</mo><mi>j</mi></munder><msub><mi>w</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><msub><mi>z</mi><mi>i</mi></msub><msub><mi>z</mi><mi>j</mi></msub></mrow>
      <mrow><munder><mo largeop="true">&#x2211;</mo><mi>i</mi></munder><msup><msub><mi>z</mi><mi>i</mi></msub><mn>2</mn></msup></mrow></mfrac>
    </mrow></math>""",
}


def _wrap_fences(mathml):
    """MML2OMML turns <mrow><mo>(</mo> a <mo>-</mo> b <mo>)</mo></mrow> into a delimiter whose
    inner operator becomes the separator character (sepChr), so "(v_ij - v_j+)" is stored as two
    delimiter cells joined by a hyphen glyph. Wrapping the inside in one <mrow> keeps the operator
    as a real operator in a single cell (audit 2026-10-09, F-3/F-4)."""
    ns = f"{{{_MML_NS}}}"
    for row in list(mathml.iter(f"{ns}mrow")):
        kids = list(row)
        if (len(kids) > 3 and kids[0].tag == f"{ns}mo" and kids[-1].tag == f"{ns}mo"
                and (kids[0].text or "").strip() == "(" and (kids[-1].text or "").strip() == ")"):
            inner = etree.SubElement(row, f"{ns}mrow")
            for k in kids[1:-1]:
                inner.append(k)
            row.remove(kids[-1])
            row.append(kids[-1])
            row.insert(1, inner)


def eq(text):
    doc, _apply_tnr = _STATE['doc'], _STATE['font']
    """Insert an editable, professionally built-up Word equation (OMML)."""
    match = re.search(r"\((\d+)\)\s*$", text)
    if not match:
        raise ValueError(f"Equation number missing from: {text}")
    number = int(match.group(1))
    mathml = etree.fromstring(_EQUATION_MATHML[number].encode("utf-8"))
    _wrap_fences(mathml)
    omml = _MML2OMML(mathml).getroot()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent       = Cm(0)
    p.paragraph_format.space_before      = Pt(3)
    p.paragraph_format.space_after       = Pt(3)
    p.paragraph_format.first_line_indent = Cm(0)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(7.96), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Cm(15.92), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    p._p.append(omml)
    p.add_run("\t")
    r = p.add_run(f"({_STATE['prefix']}{number})")
    _apply_tnr(r, 10.5)
    return p
