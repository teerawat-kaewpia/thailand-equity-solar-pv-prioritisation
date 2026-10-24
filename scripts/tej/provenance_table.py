"""One provenance file for everything the study uses: data, maps and figures, policy documents.

Built from the sources that already carry the details, so nothing is typed twice:
  data             Table S10 of the built Supplementary (same rows the paper prints)
  figure / map     Figures/FIGURES_MANIFEST.csv plus the input layers of each figure (FIG_INPUTS below)
  policy document  Data/policy_sources/SOURCE_LINKS.csv (issuer, URL, download date, sha256)
Output: Docs/PROVENANCE.csv (also copied into the release package by build_release_package.py).

Source licences: only stated where the provider publishes an open licence. Thai government statistics and
documents are marked "agency terms": no open licence is stated in the files used here.
"""
import csv
import re
import sys
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _paths import BASE
from _floats import SUPP_T

LICENCE = [  # (pattern in the dataset name, licence as published by the provider)
    (r"Global Solar Atlas", "CC BY 4.0"), (r"WorldPop", "CC BY 4.0"), (r"OpenStreetMap", "ODbL 1.0"),
    (r"Hansen", "CC BY 4.0"), (r"Copernicus DEM", "Copernicus DEM licence (free, attribution)"),
    (r"GHSL", "CC BY 4.0"), (r"GADM", "GADM licence (non-commercial use, no redistribution)"),
    (r"Open Buildings", "CC BY 4.0 or ODbL 1.0 (dual)"), (r"GridFinder", "CC BY 4.0"),
    (r"Natural Earth", "Public domain"), (r"VIIRS", "Agency terms (EOG/NOAA) and TGO agency terms"),
    (r"NESDC|NSO|Bangkok Metropolitan|Author-built", "Agency terms (Thai government statistics)"),
]
FIG_INPUTS = {
    "framework": "none (method diagram)",
    "study-region": "GADM 4.1 provinces and regions, Natural Earth Admin 0 (neighbouring countries)",
    "criteria-maps": "GADM 4.1, C1 Global Solar Atlas, C2 NESDC Table 1.5, C3 WorldPop, C4 OSM power lines, C5 Hansen + Copernicus DEM + GHSL",
    "pai-map": "GADM 4.1, province PAI (Run 4)",
    "zones": "GADM 4.1, K-means zones (Run 4)",
    "deployment": "OSM solar power plants, province PAI (Run 4)",
    "S-district": "GADM 4.1 Bangkok districts, Google Open Buildings v3, BMA district statistics",
    "corr": "province criterion matrix (Run 4)", "weights": "EWM weights (Run 4)",
    "lisa": "GADM 4.1, province PAI and LISA (Run 4)", "ranking": "province PAI (Run 4)",
    "sens": "weighting sensitivity ranks (Run 4)", "radar": "zone means of the criterion matrix (Run 4)",
}


STAGE = [  # (pattern in the S10 role, analysis stage)
    (r"^C[1-5] ", "province PAI criterion"), (r"Administrative", "aggregation units for all analyses"),
    (r"^Bangkok", "Bangkok district PAI (Supplementary S12)"), (r"Modelled MV", "allocation scenarios, MV-proximity classes"),
    (r"Solar deployment", "deployment comparison (Fig. 6)"), (r"construct check", f"construct check (Table {SUPP_T['construct']})"),
    (r"Map base layer", "map context only (Fig. 2)"), (r"LV-density", "sensitivity only (Table S9)"),
    (r"Step 1 screen","excluded at screening (post hoc diagnostic, Section S13.11)"),
]
# maps carry third-party layers whose terms are not replaced by this study's CC BY 4.0
MAPNOTE = {k: "; boundaries GADM 4.1 under GADM terms" for k in ("study-region", "criteria-maps", "pai-map", "zones", "S-district", "lisa")}
MAPNOTE["study-region"] += "; Natural Earth (public domain)"
MAPNOTE["criteria-maps"] += "; C4 derived from OpenStreetMap (ODbL, (c) OpenStreetMap contributors)"


def stage(role):
    hits = [st for pat, st in STAGE if re.search(pat, role)]
    assert len(hits) == 1, (role, hits)
    return hits[0]


def licence(name):
    for pat, lic in LICENCE:
        if re.search(pat, name):
            return lic
    return "see provider"


rows = []
sup = docx.Document(str(BASE / "Manuscript" / "DER_TEJ_Supplementary.docx"))
s10 = next(t for t in sup.tables if "Dataset and provider" in [c.text for c in t.rows[0].cells])
hdr = [c.text for c in s10.rows[0].cells]
for i, r in enumerate(s10.rows[1:], start=1):
    d = dict(zip(hdr, (c.text.strip() for c in r.cells)))
    enters = stage(d["Role"])
    rows.append({"id": f"D{i:02d}", "category": "data", "item": d["Dataset and provider"], "role_in_study": d["Role"],
                 "version_or_scale": d["Version / native scale"], "date": d["Period / access"],
                 "source_or_route": d["Access route / script"], "used_for": d["Use and processing"],
                 "analysis_stage": enters, "source_licence": licence(d["Dataset and provider"]), "sha256": "",
                 "where_documented": "Supplementary Table S10"})

with open(BASE / "Figures" / "FIGURES_MANIFEST.csv", encoding="utf-8") as fh:
    for i, f in enumerate(csv.DictReader(fh), start=1):
        rows.append({"id": f"F{i:02d}", "category": "map" if "Map" in f["file"] or "Study_Area" in f["file"] or "Criteria_Maps" in f["file"]
                     or "District" in f["file"] else "figure", "item": f"{f['label']} ({f['document']})",
                     "role_in_study": f["key"], "version_or_scale": f"PNG 600 dpi, TIFF 1000 dpi ({f['tiff_pixels']} px)",
                     "date": "Run 4", "source_or_route": f["script"], "used_for": FIG_INPUTS[f["key"]],
                     "analysis_stage": "figure", "source_licence": "CC BY 4.0 (this study)" + MAPNOTE.get(f["key"], ""), "sha256": f["sha256"],
                     "where_documented": "Figures/FIGURES_MANIFEST.csv"})

# check 2026-10-24: reference numbers as printed in the manuscript (SOURCE_LINKS converts them), an English
# title next to each Thai file name, and the current number of the zone-to-policy table (was "Table 15")
from _floats import resolve_text
with open(BASE / "Data" / "policy_sources" / "SOURCE_LINKS.csv", encoding="utf-8-sig") as fh:
    for p in csv.DictReader(fh):
        rows.append({"id": p["policy_id"], "category": "policy document",
                     "item": f"{p['title_en']} (file: {p['file']})",
                     "role_in_study": (f"cited as reference [{p['manuscript_ref']}]" if p["manuscript_ref"]
                                       else "not in the reference list, policy register only (Supplementary S11)"),
                     "version_or_scale": "",
                     "date": f"downloaded {p['download_date']}", "source_or_route": p["url"],
                     "used_for": resolve_text("zone-to-policy screen (Table {T:policy}, Supplementary S11)"),
                     "analysis_stage": "policy screen (not part of the PAI)",
                     "source_licence": f"Agency terms ({p['issuer']})", "sha256": p["sha256"],
                     "where_documented": "Data/policy_sources/SOURCE_LINKS.csv, Docs/D5_POLICY_REGISTER.md"})

# Agency web pages cited in the manuscript (check 2026-10-24: the P13 row had been added to the CSV by hand and
# was lost on regeneration, so the pages are generated here). Numbers are printed reference numbers.
import json as _json
import re as _re
import docx as _docx
_rn = {int(k): v for k, v in _json.loads((BASE / "Data" / "qa" / "ref_renumber.json").read_text(encoding="utf-8")).items()}
_refs = _json.loads((BASE / "Data" / "references.json").read_text(encoding="utf-8"))["entries"]
_ps = [q.text for q in _docx.Document(str(BASE / "Manuscript" / "DER_TEJ_Manuscript.docx")).paragraphs if q.text.strip()]
_k = max(i for i, t in enumerate(_ps) if t == "References")
_printed = {int(_re.match(r"\[(\d+)\]", t).group(1)): t for t in _ps[_k + 1:]}
WEB = [("P13", "EPPO opens public consultation on the draft Power Development Plan 2026 to 2050 (PDP2026)", "EPPO opens public consultation",
        "accessed 2026-10-16", "draft plan, planning context only (Sections 1 and 5.5)", "policy background"),
       ("P12", "MEA, Authority [web page]: the MEA service area (Bangkok, Nonthaburi, Samut Prakan)", "Authority [web page]",
        "accessed 2026-10-08", "service areas of MEA and PEA (Section 3.1)", "study region"),
       ("P14", "PEA, Business and services of PEA [web page]: the PEA service area (74 provinces)", "Business and services of PEA",
        "accessed 2026-10-09", "service areas of MEA and PEA (Section 3.1), scenario analysis set", "study region")]
for pid, item, key, acc, use, stage in WEB:
    raw = next(i for i, s in enumerate(_refs, 1) if key in s)
    num = _rn[raw]
    assert key in _printed[num], (pid, num)
    url = _re.search(r"https?://\S+", _refs[raw - 1]).group(0).rstrip(".")
    if pid == "P13":   # the full item URL, the reference prints its decoded form
        url = ("https://www.eppo.go.th/information-services/advertise/news-eppo/%E0%B8%AA%E0%B8%99%E0%B8%9E-"
               "%E0%B9%80%E0%B8%9B%E0%B8%B4%E0%B8%94%E0%B8%A3%E0%B8%B1%E0%B8%9A%E0%B8%9F%E0%B8%B1%E0%B8%87"
               "%E0%B8%84%E0%B8%A7%E0%B8%B2%E0%B8%A1%E0%B8%84%E0%B8%B4%E0%B8%94%E0%B9%80-4/")
    rows.append({"id": pid, "category": "policy web page", "item": item, "role_in_study": f"cited as reference [{num}]",
                 "version_or_scale": "", "date": acc, "source_or_route": url, "used_for": use,
                 "analysis_stage": stage, "source_licence": "Agency web page", "sha256": "",
                 # P12 and P13 have D5 register entries, the PEA page (P14, new ID) has none
                 "where_documented": ("Docs/D5_POLICY_REGISTER.md" if pid in ("P12", "P13")
                                      else "manuscript reference list (no D5 register entry)")})

assert not any("{T:" in v or "{F:" in v for r in rows for v in r.values()), "unresolved float token"
out = BASE / "Docs" / "PROVENANCE.csv"
with open(out, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerows(rows)
from collections import Counter
print(out, Counter(r["category"] for r in rows), "| data entering the PAI:",
      [r["role_in_study"] for r in rows if r["analysis_stage"] == "province PAI criterion"])
